# -*- coding: utf-8 -*-
"""テキストボックスの必要行数・必要高さを、改行(a:br)と段落を正しく数えて推定する"""
import unicodedata, math
from pptx.util import Emu
EMU = 914400
NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
SAFE = 1.00           # フォント置換ぶんの安全率（実効幅をこの割合までに収める）

def wide(c):
    return unicodedata.east_asian_width(c) in ('W', 'F', 'A')

def em_len(t):
    return sum(1.0 if wide(c) else 0.55 for c in t)

def seg_runs(pa):
    """段落を a:br で区切り、[[(text, ptsize), ...], ...] を返す"""
    segs, cur = [], []
    for ch in pa._p:
        tag = ch.tag
        if tag == NS + 'br':
            segs.append(cur); cur = []
        elif tag == NS + 'r':
            t = ''.join(n.text or '' for n in ch.findall(NS + 't'))
            sz = None
            rPr = ch.find(NS + 'rPr')
            if rPr is not None and rPr.get('sz'):
                sz = int(rPr.get('sz')) / 100.0
            cur.append((t, sz))
    segs.append(cur)
    return segs

def para_size(pa, default=18.0):
    ss = [r.font.size.pt for r in pa.runs if r.font.size]
    return max(ss) if ss else default

def flat(sh, out):
    if sh.shape_type == 6:
        for c in sh.shapes: flat(c, out)
    else:
        out.append(sh)

def metrics(sh, scale=1.0, default=18.0):
    """(必要行数, 必要高さin, 最大フォントpt) を返す"""
    tf = sh.text_frame
    W = (sh.width or 0) / EMU
    ml = (tf.margin_left if tf.margin_left is not None else 91440) / EMU
    mr = (tf.margin_right if tf.margin_right is not None else 91440) / EMU
    mt = (tf.margin_top if tf.margin_top is not None else 45720) / EMU
    mb = (tf.margin_bottom if tf.margin_bottom is not None else 45720) / EMU
    avail = (W - ml - mr) * SAFE
    total_h, lines, mx = 0.0, 0, 0.0
    for pa in tf.paragraphs:
        base = para_size(pa, default) * scale
        mx = max(mx, base)
        n = 0
        for seg in seg_runs(pa):
            if not seg:
                n += 1; continue
            w = sum(em_len(t) * ((sz or para_size(pa, default)) * scale) / 72.0 for t, sz in seg)
            n += max(1, math.ceil(w / avail - 1e-9)) if avail > 0 else 1
        ls = pa.line_spacing if isinstance(pa.line_spacing, float) else 1.0
        sb = (pa.space_before.pt if pa.space_before else 0) / 72.0
        sa = (pa.space_after.pt if pa.space_after else 0) / 72.0
        total_h += n * base / 72.0 * 1.22 * ls + sb + sa
        lines += n
    return lines, total_h + mt + mb, mx
